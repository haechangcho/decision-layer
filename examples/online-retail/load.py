"""Load the pinned UCI workbook into a separate, reproducible example database."""

from __future__ import annotations

import hashlib
import os
from pathlib import Path
import shutil
import urllib.request
from zipfile import ZipFile

import psycopg


SOURCE_URL = "https://archive.ics.uci.edu/static/public/502/online%2Bretail%2Bii.zip"
SOURCE_SHA256 = "572e36277c2390fbfde10664750731e0a86f55e33470d91919085f0408e67bfb"
WORKBOOK_SHA256 = "bcbe73b35f5b7babf197fb0cb983a11f5d9ff929078d4aa53d171b1f2df2e980"
WORKBOOK = "online_retail_II.xlsx"
EXPECTED_ROWS = 1_067_371
HEADERS = ("Invoice", "StockCode", "Description", "Quantity", "InvoiceDate", "Price", "Customer ID", "Country")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def archive_path() -> Path:
    supplied = Path(os.environ.get("RETAIL_ARCHIVE", "/input/online-retail-ii.zip"))
    cache = Path(os.environ.get("RETAIL_CACHE", "/cache"))
    archive = supplied if supplied.is_file() else cache / "online-retail-ii.zip"
    archive.parent.mkdir(parents=True, exist_ok=True)
    if not archive.exists():
        staged = archive.with_suffix(".part")
        try:
            with urllib.request.urlopen(SOURCE_URL, timeout=120) as response, staged.open("wb") as target:
                shutil.copyfileobj(response, target)
            if sha256(staged) != SOURCE_SHA256:
                raise ValueError("UCI archive checksum changed; inspect the upstream dataset before updating the pin")
            staged.replace(archive)
        finally:
            staged.unlink(missing_ok=True)
    if sha256(archive) != SOURCE_SHA256:
        raise ValueError(f"Archive checksum mismatch: {archive}. Use the pinned UCI file or update the reviewed pin.")
    return archive


def workbook_path(archive: Path) -> Path:
    cache = Path(os.environ.get("RETAIL_CACHE", "/cache"))
    cache.mkdir(parents=True, exist_ok=True)
    workbook = cache / WORKBOOK
    if workbook.exists() and sha256(workbook) == WORKBOOK_SHA256:
        return workbook
    with ZipFile(archive) as source:
        if source.namelist() != [WORKBOOK]:
            raise ValueError("Unexpected UCI archive contents")
        with source.open(WORKBOOK) as input_file, workbook.open("wb") as output_file:
            shutil.copyfileobj(input_file, output_file)
    if sha256(workbook) != WORKBOOK_SHA256:
        workbook.unlink(missing_ok=True)
        raise ValueError("Workbook checksum mismatch")
    return workbook


def identifier(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def normalized_row(values: tuple[object, ...], line_id: int) -> tuple[object, ...]:
    invoice, stock, description, quantity, invoice_at, price, customer, country = values
    if invoice is None or stock is None or quantity is None or invoice_at is None or price is None:
        raise ValueError(f"Required source field missing at line {line_id}")
    return (line_id, identifier(invoice), identifier(stock), description, int(quantity), invoice_at,
            str(price), identifier(customer), country)


def load() -> None:
    from openpyxl import load_workbook

    workbook = workbook_path(archive_path())
    with psycopg.connect(os.environ["DATABASE_URL"]) as connection:
        with connection.cursor() as cursor:
            cursor.execute("CREATE SCHEMA IF NOT EXISTS retail")
            cursor.execute("""CREATE TABLE IF NOT EXISTS retail.invoice_line (
                line_id bigint PRIMARY KEY,
                invoice_no text NOT NULL,
                stock_code text NOT NULL,
                description text,
                quantity integer NOT NULL,
                invoice_at timestamp NOT NULL,
                unit_price numeric(14, 4) NOT NULL,
                customer_id text,
                country text
            )""")
            cursor.execute("""CREATE TABLE IF NOT EXISTS retail.import_state (
                source_sha256 text PRIMARY KEY,
                row_count integer NOT NULL,
                loaded_at timestamptz NOT NULL DEFAULT now()
            )""")
            cursor.execute("""DO $$ BEGIN
                IF NOT EXISTS (SELECT 1 FROM pg_roles WHERE rolname = 'retail_reader') THEN
                    CREATE ROLE retail_reader LOGIN PASSWORD 'local-read-only';
                END IF;
            END $$""")
            cursor.execute("GRANT USAGE ON SCHEMA retail TO retail_reader")
            cursor.execute("GRANT SELECT ON retail.invoice_line TO retail_reader")
            cursor.execute("SELECT row_count FROM retail.import_state WHERE source_sha256 = %s", (SOURCE_SHA256,))
            imported = cursor.fetchone()
            if imported and imported[0] == EXPECTED_ROWS:
                cursor.execute("SELECT count(*) FROM retail.invoice_line")
                if cursor.fetchone()[0] == EXPECTED_ROWS:
                    print(f"Pinned UCI dataset already loaded: {EXPECTED_ROWS:,} rows", flush=True)
                    return

            cursor.execute("TRUNCATE retail.invoice_line")
            cursor.execute("DELETE FROM retail.import_state")
            source = load_workbook(workbook, read_only=True, data_only=True)
            line_id = 0
            try:
                with cursor.copy("""COPY retail.invoice_line
                    (line_id, invoice_no, stock_code, description, quantity, invoice_at, unit_price, customer_id, country)
                    FROM STDIN""") as copy:
                    for sheet in source.worksheets:
                        rows = sheet.iter_rows(values_only=True)
                        if tuple(next(rows)) != HEADERS:
                            raise ValueError(f"Unexpected UCI columns in {sheet.title}")
                        for row in rows:
                            line_id += 1
                            copy.write_row(normalized_row(row, line_id))
                            if line_id % 100_000 == 0:
                                print(f"Loaded {line_id:,} rows", flush=True)
            finally:
                source.close()
            if line_id != EXPECTED_ROWS:
                raise ValueError(f"Expected {EXPECTED_ROWS:,} rows, found {line_id:,}")
            cursor.execute("CREATE INDEX IF NOT EXISTS invoice_line_at_idx ON retail.invoice_line (invoice_at)")
            cursor.execute("CREATE INDEX IF NOT EXISTS invoice_line_country_idx ON retail.invoice_line (country)")
            cursor.execute("CREATE INDEX IF NOT EXISTS invoice_line_stock_idx ON retail.invoice_line (stock_code)")
            cursor.execute("INSERT INTO retail.import_state (source_sha256, row_count) VALUES (%s, %s)",
                           (SOURCE_SHA256, line_id))
            cursor.execute("ANALYZE retail.invoice_line")
    print(f"UCI Online Retail II loaded: {line_id:,} rows", flush=True)


if __name__ == "__main__":
    load()
